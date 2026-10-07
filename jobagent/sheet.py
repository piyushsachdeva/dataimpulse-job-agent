"""Spreadsheet of every job the agent has seen: jobs.xlsx (and jobs.csv).

SQLite (jobs.db) is the source of truth; the sheet is regenerated from it after every
search or assessment, so it is always current. Jobs first seen today are highlighted.
"""

import csv
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import store

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "jobs.xlsx"
CSV = ROOT / "jobs.csv"

COLUMNS = [  # (header, db field, width)
    ("First seen", "first_seen", 18),
    ("Company", "company", 16),
    ("Title", "title", 42),
    ("Location", "location", 28),
    ("Country", "country", 16),
    ("Sponsorship", "sponsorship", 12),
    ("AI fit score", "score", 11),
    ("AI notes", "notes", 50),
    ("Status", "status", 12),
    ("URL", "url", 50),
]
NEW = PatternFill("solid", fgColor="FFF4CC")
SPONSOR = {"yes": "C6EFCE", "no": "FFC7CE"}


def export() -> str:
    rows = store.sheet_rows()
    today = datetime.now(timezone.utc).date().isoformat()  # first_seen is stored in UTC

    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append([h for h, _, _ in COLUMNS])
    for r in rows:
        ws.append([r.get(f) for _, f, _ in COLUMNS])
        row = ws.max_row
        if str(r.get("first_seen", "")).startswith(today):
            for cell in ws[row]:
                cell.fill = NEW
        if r.get("sponsorship") in SPONSOR:
            ws.cell(row, 6).fill = PatternFill("solid", fgColor=SPONSOR[r["sponsorship"]])
        url = ws.cell(row, len(COLUMNS))
        if url.value:
            url.hyperlink, url.style = url.value, "Hyperlink"

    for i, (_, _, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):  # AI notes
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    if rows:
        tab = Table(displayName="Jobs", ref=f"A1:{ws.cell(ws.max_row, len(COLUMNS)).coordinate}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tab)
    wb.save(XLSX)

    with open(CSV, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow([h for h, _, _ in COLUMNS])
        w.writerows([[r.get(f) for _, f, _ in COLUMNS] for r in rows])

    new_today = sum(str(r.get("first_seen", "")).startswith(today) for r in rows)
    return f"{XLSX.name}: {len(rows)} jobs ({new_today} new today, highlighted)"
